"""Deterministic PDF, DOCX, and text parsing with OCR for scanned PDF pages."""
from io import BytesIO
from pathlib import Path
from typing import Any

try:
    import pymupdf
except ImportError:  # pragma: no cover - compatibility with older PyMuPDF releases
    import fitz as pymupdf

from docx import Document


class DocumentParseError(ValueError):
    """Structured error raised when a source cannot be parsed safely."""
    def __init__(self, code: str, message: str):
        self.code = code
        super().__init__(f"{code}: {message}")


def _result(file_name: str, file_format: str, pages: list[dict[str, Any]]) -> dict[str, Any]:
    text = "\n\n".join(page["text"] for page in pages if page["text"])
    if not text.strip():
        raise DocumentParseError("NO_EXTRACTABLE_TEXT", f"{file_name} contains no extractable text, including OCR")
    return {
        "fileName": file_name,
        "format": file_format,
        "pages": pages,
        "text": text,
        "charCount": len(text),
    }


def _ocr_page(page: Any, page_number: int) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as error:
        raise DocumentParseError(
            "OCR_UNAVAILABLE",
            "scanned PDF detected but pytesseract/Pillow are not installed; install Tesseract OCR and pytesseract",
        ) from error
    try:
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
        image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        text = pytesseract.image_to_string(image, timeout=15)
        return text.strip()
    except (ImportError, OSError, RuntimeError, ValueError) as error:
        raise DocumentParseError(
            "OCR_FAILED", f"OCR failed on PDF page {page_number}: {type(error).__name__}"
        ) from error


def _parse_pdf(data: bytes, file_name: str) -> dict[str, Any]:
    try:
        document = pymupdf.open(stream=data, filetype="pdf")
    except (pymupdf.FileDataError, pymupdf.EmptyFileError, OSError, RuntimeError, ValueError) as error:
        raise DocumentParseError("INVALID_DOCUMENT", f"could not open PDF: {error}") from error
    try:
        pages = []
        for number, page in enumerate(document, start=1):
            text = (page.get_text("text") or "").strip()
            if not text:
                text = _ocr_page(page, number)
            if text:
                pages.append({"page": number, "text": text})
    finally:
        document.close()
    return _result(file_name, "pdf", pages)


def _parse_docx(data: bytes, file_name: str) -> dict[str, Any]:
    try:
        document = Document(BytesIO(data))
    except (OSError, RuntimeError, ValueError, KeyError) as error:
        raise DocumentParseError("INVALID_DOCUMENT", f"could not open DOCX: {error}") from error
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return _result(file_name, "docx", [{"page": None, "text": "\n".join(parts)}])


def _parse_text(data: bytes, file_name: str, file_format: str) -> dict[str, Any]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise DocumentParseError("INVALID_DOCUMENT", f"text file is not valid UTF-8: {error}") from error
    return _result(file_name, file_format, [{"page": None, "text": text}])


def parse_document(data: bytes, file_name: str, mime_type: str | None = None) -> dict[str, Any]:
    """Parse supported source bytes and OCR scanned PDF pages without fabricating text."""
    if not isinstance(data, (bytes, bytearray)) or not data:
        raise DocumentParseError("NO_EXTRACTABLE_TEXT", f"{file_name} is empty")
    suffix = Path(file_name).suffix.lower()
    if suffix == ".pdf" or mime_type == "application/pdf":
        return _parse_pdf(bytes(data), file_name)
    if suffix == ".docx" or mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _parse_docx(bytes(data), file_name)
    if suffix in {".md", ".markdown"}:
        return _parse_text(bytes(data), file_name, "markdown")
    if suffix == ".txt" or mime_type == "text/plain":
        return _parse_text(bytes(data), file_name, "text")
    raise DocumentParseError("UNSUPPORTED_FORMAT", f"unsupported source format for {file_name}")
